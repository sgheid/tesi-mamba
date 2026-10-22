import io
import math
import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .formato import FORMATI

BLU_SCURO = "1F3864"
SFONDO_INTESTAZIONE = PatternFill("solid", fgColor=BLU_SCURO)
SFONDO_EVIDENZA = PatternFill("solid", fgColor="FFF2CC")
SFONDO_SOTTOFASCIA = PatternFill("solid", fgColor="DDEBF7")
FONT_INTESTAZIONE = Font(bold=True, color="FFFFFF")
FONT_SEZIONE = Font(bold=True, size=12, color=BLU_SCURO)
BORDO = Border(bottom=Side(style="thin", color="BFBFBF"))
SINISTRA = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTRO_VERTICALE = Alignment(horizontal="left", vertical="center")


def _nome_modello(info):
    return info.get("modello", "modello").split("/")[-1].split(" @ ")[0]


def nome_file(ris):
    data = ris.info.get("data", datetime.now())
    # I parametri nel nome (es. _FP16): esecuzioni della stessa fase nello stesso secondo non si sovrascrivono
    parametri = "".join("_" + re.sub(r"\W+", "-", str(v)).strip("-") for v in ris.info.get("parametri", {}).values())
    return f"{data:%Y%m%d_%H%M%S}_{_nome_modello(ris.info)}_Fase{ris.codice}{parametri}.xlsx"


def _nome_foglio(testo, usati):
    base = re.sub(r"[\\/*?:\[\]]", "", testo)[:28] or "Dati"
    nome, k = base, 2
    while nome in usati:
        nome, k = f"{base[:25]} {k}", k + 1
    usati.add(nome)
    return nome


def _valore_cella(v):
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v.item() if hasattr(v, "item") else v


def _intestazione(cella):
    cella.font = FONT_INTESTAZIONE
    cella.fill = SFONDO_INTESTAZIONE


# Mattoni comuni ai due file: ognuno scrive da una riga e restituisce la riga successiva
def _sezione(ws, riga, titolo):
    ws.cell(riga, 1, titolo).font = FONT_SEZIONE
    return riga + 1


# Etichetta nelle colonne A-B, valore nelle colonne da C a ultima
def _coppia(ws, riga, etichetta, valore, ultima):
    ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=2)
    ws.merge_cells(start_row=riga, start_column=3, end_row=riga, end_column=ultima)
    e = ws.cell(riga, 1, etichetta)
    e.font = Font(bold=True)
    e.alignment = SINISTRA
    ws.cell(riga, 3, valore).alignment = SINISTRA
    return riga + 1


# Il testo analizzato in una riga alta abbastanza da leggerlo (Excel non adatta da solo le celle unite)
def _testo_analizzato(ws, riga, testo, ultima):
    caratteri_per_riga = 23 * (ultima - 2)
    n_righe = sum(len(p) // caratteri_per_riga + 1 for p in testo.split("\n"))
    ws.row_dimensions[riga].height = min(409, 15 * (n_righe + 1))
    return _coppia(ws, riga, "Testo analizzato", testo[:32000], ultima)


def _info_modello(info):
    return [
        ("Modello", info.get("modello", "")),
        ("Cartella", info.get("percorso", "")),
        ("Commit dei pesi", info.get("commit", "")),
        ("Struttura", info.get("struttura", "")),
        ("Dispositivo", info.get("dispositivo", "")),
    ]


def _info_testo(info):
    if info.get("testo") is None:
        return []
    return [("Fonte del testo", info.get("fonte_testo", "")), ("Token analizzati", str(info.get("n_token", "")))]


def _sintesi(ws, riga, ris, ultima):
    for ind in ris.indicatori:
        riga = _coppia(ws, riga, ind.etichetta, ind.valore + (f"  -  {ind.nota}" if ind.nota else ""), ultima)
    return riga


# Tabella per strato, con le righe evidenziate e la nota a destra; restituisce anche la riga d'intestazione
def _tabella(ws, riga, ris):
    col = ris.colonne
    for j, c in enumerate(col, start=1):
        cella = ws.cell(riga, j, f"{c.titolo}\n{c.sottotitolo}".strip())
        _intestazione(cella)
        cella.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[riga].height = 32
    riga_intestazione = riga
    riga += 1
    for r in ris.righe:
        evid = r.get("strato") in ris.evidenzia
        for j, c in enumerate(col, start=1):
            cella = ws.cell(riga, j, _valore_cella(r.get(c.chiave)))
            cella.number_format = FORMATI[c.formato][1]
            cella.border = BORDO
            if c.chiave == "strato":
                cella.alignment = Alignment(horizontal="center")
            if evid:
                cella.fill = SFONDO_EVIDENZA
        if evid:
            nota = ws.cell(riga, len(col) + 1, ris.evidenzia[r["strato"]])
            nota.font = Font(italic=True)
            nota.alignment = Alignment(indent=1)
        riga += 1
    return riga_intestazione, riga


def _legenda(ws, riga, voci, ultima):
    for nome_col, spiegazione in voci:
        riga = _coppia(ws, riga, nome_col, spiegazione, ultima)
    return riga


def _larghezze(ws, ultima):
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    for j in range(3, ultima + 1):
        ws.column_dimensions[get_column_letter(j)].width = 20


def _stampa_orizzontale(ws):
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _in_byte(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# Excel di un Risultato: foglio della fase (info, sintesi, tabella, legenda) + un foglio per grafico
def crea_excel(ris):
    wb = openpyxl.Workbook()
    ws = wb.active
    usati = set()
    ws.title = _nome_foglio(f"Fase {ris.codice}", usati)

    n_col = len(ris.colonne)
    ultima = max(n_col + 1, 7)
    info = ris.info

    ws.cell(1, 1, ris.titolo).font = Font(bold=True, size=14)
    ws.cell(2, 1, ris.descrizione).font = Font(italic=True)
    riga = _sezione(ws, 4, "INFORMAZIONI")
    elenco = _info_modello(info) + [
        ("Formato numerico", ris.note_formato),
        ("Data e ora", f"{info.get('data', datetime.now()):%d/%m/%Y %H:%M:%S}"),
        ("Durata", f"{info.get('durata', 0):.1f} s"),
    ]
    elenco += [(f"Parametro: {k}", str(v)) for k, v in info.get("parametri", {}).items()]
    elenco += _info_testo(info)
    for etichetta, valore in elenco:
        riga = _coppia(ws, riga, etichetta, valore, ultima)
    if info.get("testo") is not None:
        riga = _testo_analizzato(ws, riga, info["testo"], ultima)
    riga += 1

    riga = _sezione(ws, riga, "SINTESI")
    riga = _sintesi(ws, riga, ris, ultima) + 1

    riga = _sezione(ws, riga, "RISULTATI PER STRATO")
    riga_intestazione, riga = _tabella(ws, riga, ris)
    riga += 1

    riga = _sezione(ws, riga, "COME LEGGERE LE COLONNE")
    _legenda(ws, riga, ris.legenda, ultima)

    _larghezze(ws, ultima)
    ws.column_dimensions[get_column_letter(n_col + 1)].width = 26
    ws.freeze_panes = ws.cell(riga_intestazione + 1, 1)
    _stampa_orizzontale(ws)

    # Un foglio con i dati di ogni grafico
    for g in ris.grafici:
        if g.dati is None or len(g.dati) == 0:
            continue
        wg = wb.create_sheet(_nome_foglio(g.titolo, usati))
        wg.cell(1, 1, g.titolo).font = Font(bold=True, size=12)
        if g.spiegazione:
            wg.cell(2, 1, g.spiegazione).font = Font(italic=True)
        for j, nome in enumerate(g.dati.columns, start=1):
            _intestazione(wg.cell(4, j, str(nome)))
            wg.column_dimensions[get_column_letter(j)].width = max(14, len(str(nome)) + 2)
        for i, riga_dati in enumerate(g.dati.itertuples(index=False), start=5):
            for j, v in enumerate(riga_dati, start=1):
                wg.cell(i, j, _valore_cella(v))
        wg.freeze_panes = "A5"

    return _in_byte(wb)


def salva_excel(ris, cartella="risultati"):
    os.makedirs(cartella, exist_ok=True)
    percorso = os.path.join(cartella, nome_file(ris))
    with open(percorso, "wb") as f:
        f.write(crea_excel(ris))
    return percorso


# File unico: tutte le analisi in un solo foglio, senza i dati dei grafici.
# analisi = [(fase, parametri, risultato o None, (messaggio, dettaglio) o None)] nell'ordine di esecuzione
ULTIMA_COPPIA = 9


def _riusciti(analisi):
    return [ris for _, _, ris, _ in analisi if ris is not None]


def nome_file_unico(analisi):
    riusciti = _riusciti(analisi)
    data = min(r.info.get("data", datetime.now()) for r in riusciti)
    return f"{data:%Y%m%d_%H%M%S}_{_nome_modello(riusciti[0].info)}_Tutte-le-fasi.xlsx"


def _parametri(parametri):
    return " · ".join(str(v) for v in parametri.values()) or "-"


# Una riga colorata con il testo unito sulle colonne da 1 a larghezza
def _fascia(ws, riga, testo, larghezza, sfondo, font, altezza):
    ws.merge_cells(start_row=riga, start_column=1, end_row=riga, end_column=larghezza)
    for j in range(1, larghezza + 1):
        ws.cell(riga, j).fill = sfondo
    cella = ws.cell(riga, 1, testo)
    cella.font = font
    cella.alignment = CENTRO_VERTICALE
    ws.row_dimensions[riga].height = altezza
    return riga + 1


def crea_excel_unico(analisi):
    riusciti = _riusciti(analisi)
    if not riusciti:
        raise ValueError("Nessuna analisi riuscita da salvare")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risultati"
    larghezza = max([len(r.colonne) + 1 for r in riusciti] + [ULTIMA_COPPIA])
    info = riusciti[0].info
    con_testo = next((r.info for r in riusciti if r.info.get("testo") is not None), None)
    n_errori = len(analisi) - len(riusciti)

    ws.cell(1, 1, f"Analisi dei modi di Mamba — {info.get('modello', '')}").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"{len(analisi)} analisi" + (f" · {n_errori} non riuscit{'a' if n_errori == 1 else 'e'}"
                                                if n_errori else "")
            + " · sintesi e tabelle per strato; i dati dei grafici sono nell'Excel della singola fase").font = \
        Font(italic=True)

    riga = _sezione(ws, 4, "INFORMAZIONI COMUNI")
    elenco = _info_modello(info) + [
        ("Data e ora", f"{min(r.info.get('data', datetime.now()) for r in riusciti):%d/%m/%Y %H:%M:%S}"),
        ("Durata totale", f"{sum(r.info.get('durata', 0) for r in riusciti):.1f} s"),
    ]
    if con_testo:
        elenco += _info_testo(con_testo)
    for etichetta, valore in elenco:
        riga = _coppia(ws, riga, etichetta, valore, ULTIMA_COPPIA)
    if con_testo:
        riga = _testo_analizzato(ws, riga, con_testo["testo"], ULTIMA_COPPIA)
    riga += 1

    # Indice: righe riservate ora, collegamenti scritti quando si conoscono le righe delle sezioni
    riga = _sezione(ws, riga, "INDICE")
    for j, titolo in enumerate(["Fase", "Analisi", "Confronto", "Esito", "Durata", "Vai alla sezione"], start=1):
        _intestazione(ws.cell(riga, j, titolo))
    riga_indice = riga + 1
    riga = riga_indice + len(analisi) + 1

    partenze = []
    for codice in dict.fromkeys(fase.codice for fase, _, _, _ in analisi):
        gruppo = [a for a in analisi if a[0].codice == codice]
        fase = gruppo[0][0]
        inizio = riga
        riga = _fascia(ws, riga, f"FASE {fase.codice} · {fase.nome}", larghezza, SFONDO_INTESTAZIONE,
                       Font(bold=True, size=13, color="FFFFFF"), 24)
        ws.cell(riga, 1, fase.breve).font = Font(italic=True)
        riga += 2
        legenda = {}
        for _, parametri, ris, errore in gruppo:
            partenze.append(riga)
            riga = _fascia(ws, riga, _parametri(parametri), larghezza, SFONDO_SOTTOFASCIA,
                           Font(bold=True, size=12, color=BLU_SCURO), 20)
            if ris is None:
                ws.cell(riga, 1, f"Non riuscita: {errore[0]}").font = Font(bold=True, color="C00000")
                riga += 2
                continue
            i = ris.info
            ws.cell(riga, 1, f"{ris.descrizione}  Formato: {ris.note_formato} · durata {i.get('durata', 0):.1f} s · "
                             f"eseguita alle {i.get('data', datetime.now()):%H:%M:%S}").font = Font(italic=True)
            riga += 2
            riga = _sezione(ws, riga, "SINTESI")
            riga = _sintesi(ws, riga, ris, ULTIMA_COPPIA) + 1
            riga = _sezione(ws, riga, "RISULTATI PER STRATO")
            riga = _tabella(ws, riga, ris)[1] + 1
            for voce in ris.legenda:
                legenda.setdefault(voce, None)
        if legenda:
            riga = _sezione(ws, riga, f"COME LEGGERE LE COLONNE (FASE {fase.codice})")
            riga = _legenda(ws, riga, legenda, ULTIMA_COPPIA)
        # La fase si può comprimere in Excel lasciando visibile la sua fascia
        for r in range(inizio + 1, riga):
            ws.row_dimensions[r].outline_level = 1
        riga += 2

    for k, ((fase, parametri, ris, errore), partenza) in enumerate(zip(analisi, partenze)):
        r = riga_indice + k
        valori = [fase.codice, fase.nome, _parametri(parametri), "ok" if ris is not None else "non riuscita",
                  f"{ris.info.get('durata', 0):.1f} s" if ris is not None else "-"]
        for j, v in enumerate(valori, start=1):
            ws.cell(r, j, v).border = BORDO
        ws.cell(r, 1).alignment = Alignment(horizontal="center")
        if ris is None:
            ws.cell(r, 4).font = Font(color="C00000")
        vai = ws.cell(r, 6, f"→ riga {partenza}")
        vai.hyperlink = Hyperlink(ref=vai.coordinate, location=f"'Risultati'!A{partenza}")
        vai.font = Font(color="0563C1", underline="single")
        vai.border = BORDO

    _larghezze(ws, larghezza)
    ws.sheet_properties.outlinePr.summaryBelow = False
    _stampa_orizzontale(ws)
    return _in_byte(wb)


def salva_excel_unico(analisi, cartella="risultati"):
    os.makedirs(cartella, exist_ok=True)
    percorso = os.path.join(cartella, nome_file_unico(analisi))
    with open(percorso, "wb") as f:
        f.write(crea_excel_unico(analisi))
    return percorso
